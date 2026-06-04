"""把幼儿园招生名录数据导出为 Excel，方便排序/筛选/打印。"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "outputs/kindergarten/data/kindergarten_rows.json"
OUT_XLSX = ROOT / "outputs/kindergarten/海淀幼儿园2026招生名录.xlsx"


HEADER_FILL = PatternFill(start_color="D9EAF7", end_color="D9EAF7", fill_type="solid")
HEADER_FONT = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def main() -> None:
    payload = json.loads(INPUT.read_text(encoding="utf-8"))
    rows = payload["rows"]

    wb = Workbook()
    main_ws = wb.active
    main_ws.title = "园所明细"
    headers = [
        "序号",
        "街道",
        "园所名称",
        "园所性质",
        "地址信息",
        "招生咨询联系人",
        "招生咨询时间",
    ]
    main_ws.append(headers)
    for row in rows:
        main_ws.append([row.get(h, "") for h in headers])

    widths = [6, 16, 36, 18, 44, 40, 34]
    for idx, width in enumerate(widths, 1):
        main_ws.column_dimensions[get_column_letter(idx)].width = width
    for cell in main_ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
    for row in main_ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    main_ws.freeze_panes = "A2"

    summary = wb.create_sheet("街道汇总")
    summary.append(["街道", "园所数量", "公办(教育部门办)", "公办(单位办)", "民办(普惠)", "民办(非普惠)"])
    streets: dict[str, dict[str, int]] = {}
    for row in rows:
        bucket = streets.setdefault(
            row["街道"],
            {"公办(教育部门办)": 0, "公办(单位办)": 0, "民办(普惠)": 0, "民办(非普惠)": 0},
        )
        bucket[row["园所性质"]] = bucket.get(row["园所性质"], 0) + 1
    for street, bucket in sorted(streets.items(), key=lambda x: -sum(x[1].values())):
        total = sum(bucket.values())
        summary.append(
            [
                street,
                total,
                bucket.get("公办(教育部门办)", 0),
                bucket.get("公办(单位办)", 0),
                bucket.get("民办(普惠)", 0),
                bucket.get("民办(非普惠)", 0),
            ]
        )
    for idx, width in enumerate([16, 10, 18, 16, 14, 16], 1):
        summary.column_dimensions[get_column_letter(idx)].width = width
    for cell in summary[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    info = wb.create_sheet("数据来源")
    info.append(["字段", "内容"])
    info.append(["标题", payload["title"]])
    info.append(["发布机构", payload["issued_by"]])
    info.append(["发布日期", payload["issued_at"]])
    info.append(["来源 URL", payload["source_url"]])
    info.append(["园所总数", len(rows)])
    info.append(["生成脚本", "scripts/build_kindergarten_workbook.py"])
    info.column_dimensions["A"].width = 14
    info.column_dimensions["B"].width = 70
    for cell in info[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for row in info.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP

    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)
    print(f"wrote {OUT_XLSX} ({len(rows)} rows)")


if __name__ == "__main__":
    main()
